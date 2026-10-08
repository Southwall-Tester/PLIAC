"""Document ingestion contracts: real parsing, grounded extraction, isolated draft writes."""
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from learning_agent.documents import DocumentStore, atomic_json, cached_pdf_outline, chunks, document_outline, extract_local, extract_llm, hierarchy, term_id
from learning_agent.course_graph import CourseGraphStore, CourseGraphError


class DocumentsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = DocumentStore(Path(self.tmp.name) / "documents")
        self.course = CourseGraphStore(output_dir=Path(self.tmp.name) / "course")

    def tearDown(self):
        self.store.executor.shutdown(wait=True)
        self.tmp.cleanup()

    def book(self, text="知识图谱包括实体和关系。实体对齐用于知识融合。知识图谱可以支持关系抽取。", filename="教材.txt"):
        job = self.store.create(filename)
        self.store.source(job["id"]).write_text(text, encoding="utf-8")
        self.store.run(job["id"])
        return job["id"]

    def test_chunks_preserve_original_offsets_and_coverage(self):
        text = ("知识图谱包括实体。\n" * 700)
        seen = set()
        for offset, value in chunks(text):
            self.assertEqual(text[offset:offset + len(value)], value)
            self.assertLessEqual(len(value), 3200)
            seen.update(range(offset, offset + len(value)))
        self.assertEqual(len(seen), len(text))

    def test_local_provenance_is_exact_and_never_creates_prerequisite(self):
        ident = self.book()
        self.assertEqual(self.store.status(ident)["status"], "completed")
        graph, page = self.store.graph(ident), self.store.page(ident, 1)
        self.assertGreater(len(graph["nodes"]), 2)
        self.assertTrue(graph["edges"])
        self.assertNotIn("prerequisite", {e["type"] for e in graph["edges"]})
        for item in graph["nodes"] + graph["edges"]:
            for ev in item["evidence"]:
                self.assertEqual(page["text"][ev["start"]:ev["end"]], ev["quote"])

    def test_negation_and_distant_terms_do_not_become_containment(self):
        negative = extract_local("知识图谱不包含神经网络。", 1)
        distant = extract_local("知识图谱的研究定义涉及多种方法，其中包括神经网络。", 1)
        self.assertFalse(any(e["type"] == "contains" for e in negative["edges"] + distant["edges"]))
        positive = extract_local("知识图谱包括神经网络。", 1)
        self.assertTrue(any(e["type"] == "contains" for e in positive["edges"]))

    def test_terms_preserve_compounds_and_reject_ocr_fragments(self):
        extracted = extract_local("知识图谱使用远程监督进行关系抽取。实体对齐包括硬对齐和软对齐。lle1 e211", 1)
        names = {n["title"] for n in extracted["nodes"]}
        self.assertTrue({"知识图谱", "远程监督", "关系抽取", "实体对齐"} <= names)
        self.assertFalse({"lle1", "e211"} & names)

    def test_local_terms_prefer_marked_compounds_and_reject_translation_fragments(self):
        text = ("**卷积神经网络 (Convolutional Neural Network，简称 CNN)**处理图像。"
                "**支持向量机**用于分类。含义是“Non-deterministic Polynomial”。"
                "公式\\text{CNN}用于说明。家族、答案、学界只是讲解用语。"
                "**“随机化 (randomization)”**用于计算。")
        result = extract_local(text, 3, 80)
        names = {n["title"] for n in result["nodes"]}
        self.assertTrue({"卷积神经网络", "支持向量机", "随机化", "CNN"} <= names)
        self.assertFalse({"Convolutional", "Neural", "Network", "Polynomial", "text", "家族", "答案", "学界", "“随机化"} & names)
        for item in result["nodes"] + result["edges"]:
            for ev in item["evidence"]:
                self.assertEqual(text[ev["start"] - 80:ev["end"] - 80], ev["quote"])

    def test_local_defined_initial_is_preserved_without_matching_larger_words(self):
        text = "**P** 与 **NP** 是两类判定问题。公式\\text{P}与\\text{NP}用来表示类别。"
        result = extract_local(text, 1)
        by_id = {n["id"]: n["title"] for n in result["nodes"]}
        self.assertTrue({"P", "NP"} <= set(by_id.values()))
        self.assertNotIn("text", by_id.values())
        self.assertTrue(any({by_id[e["source"]], by_id[e["target"]]} == {"P", "NP"} for e in result["edges"]))
        self.assertNotIn("P", {n["title"] for n in extract_local("Polynomial time 和 NP 不是单字母定义。", 1)["nodes"]})
        self.assertTrue(all(e["type"] in {"contains", "cooccurs"} for e in result["edges"]))

    def test_local_cleanup_keeps_english_material_and_does_not_change_llm_filter(self):
        names = {n["title"] for n in extract_local("Neural networks include neurons. Networks learn representations.", 1)["nodes"]}
        self.assertTrue({"networks", "neurons"} <= names)
        reply = {"nodes": [{"id": "x", "title": "time", "quote": "time is a quantity."}], "edges": []}
        with patch("learning_agent.llm.call_llm_json", return_value=json.dumps(reply)):
            result = extract_llm("time is a quantity.", 1, config={})
        self.assertEqual([n["title"] for n in result["nodes"]], ["time"])

    def test_pdf_page_range_and_toc_are_real(self):
        import fitz
        job = self.store.create("book.pdf", start_page=2, end_page=2)
        with fitz.open() as doc:
            for text in ["Introduction text and source concepts.", "Networks\nNeurons\nNeural networks include neurons. Neural networks learn representations."]:
                page = doc.new_page()
                page.insert_text((40, 60), text)
            doc.set_toc([[1, "Introduction", 1], [1, "Networks", 2], [2, "Neurons", 2]])
            doc.save(self.store.source(job["id"]))
        self.store.run(job["id"])
        self.assertEqual(self.store.status(job["id"])["stats"]["pages"], 1)
        graph = self.store.graph(job["id"])
        self.assertTrue(graph["nodes"])
        self.assertTrue(all(ev["page"] == 2 for n in graph["nodes"] for ev in n["evidence"]))
        titles = {n["title"] for n in graph["hierarchy"]["nodes"]}
        self.assertTrue({"Networks", "Neurons"} <= titles)
        self.assertNotIn("Introduction", titles)

    def test_cached_pdf_headings_preserve_nested_structure_and_exact_crlf_evidence(self):
        header = "An Example Textbook With An Ordinary Repeated English Running Header " * 2
        pages = [
            {"page": 10, "text": f"  第1章\r\n{header}\r\n模型训练\r\n1.1 训练数据\r\n训练集用于拟合。\r\n", "ocr": True},
            {"page": 11, "text": "2\n第1章\n1.1.1 数据采样\n神经网络使用训练集。\n1.2 模型评估\n测试集用于评估。", "ocr": True},
        ]
        outline = cached_pdf_outline(pages)
        self.assertEqual([(h["level"], h["title"]) for h in outline], [
            (1, "第1章 模型训练"), (2, "1.1 训练数据"), (3, "1.1.1 数据采样"), (2, "1.2 模型评估"),
        ])
        by_page = {page["page"]: page["text"] for page in pages}
        for heading in outline:
            self.assertEqual(by_page[heading["page"]][heading["start"]:heading["end"]], heading["quote"])
        concepts = extract_local(pages[1]["text"], 11)["nodes"]
        tree = hierarchy({"id": "example", "title": "教材", "start_page": 10, "page_kind": "pdf"}, outline, concepts)
        by_title = {node["title"]: node for node in tree["nodes"]}
        self.assertEqual(by_title["1.1.1 数据采样"]["parent_id"], by_title["1.1 训练数据"]["id"])
        self.assertEqual(by_title["1.2 模型评估"]["parent_id"], by_title["第1章 模型训练"]["id"])
        for node in tree["nodes"]:
            if node["kind"] == "chapter":
                self.assertEqual(by_page[node["page"]][node["start"]:node["end"]], node["quote"])

    def test_cached_pdf_headings_reject_contents_headers_formulas_references_and_exercises(self):
        pages = [
            {"page": 1, "text": "目录\n第1章 错误目录\n1.1 目录条目……10\n1.2 目录条目……12\n1.3 目录条目……14"},
            {"page": 2, "text": "第1章\n1.1~1.12\n1.13（历史）\n2.1~2.4"},
            {"page": 10, "text": "第1章 模型训练\n1.1 训练数据\n训练集用于拟合。"},
            {"page": 11, "text": "2\n第1章\n1.2×时钟周期\n1.3 GHz\n1.4节问题讨论：可以有多种答案。\n1.5 [2] <1.1>请列举\n1.6.1 缺失父编号\n1.2 模型评估\n测试集用于评估。"},
            {"page": 12, "text": "3\n第1章\n1.3 练习\n1.1 [2] <1.1>请列举三种类型。\n1.4 2004年发布的处理器\n1.4.1 求出功耗"},
            {"page": 13, "text": "4\n第1章\n1.9 更多练习题\n1.9.1 假设处理器工作电压"},
            {"page": 14, "text": "第2章 模型部署\n2.1 推理服务\n神经网络提供服务。"},
        ]
        outline = cached_pdf_outline(pages)
        self.assertEqual([h["title"] for h in outline], [
            "第1章 模型训练", "1.1 训练数据", "1.2 模型评估", "1.3 练习", "第2章 模型部署", "2.1 推理服务",
        ])
        self.assertTrue(all(h["page"] >= 10 for h in outline))

    def test_cached_pdf_headings_do_not_mistake_numeric_tables_for_contents(self):
        table = "\n".join(["1.00", "2.37", "2.13", "1.38 1.47", "42"] * 8)
        pages = [{"page": 1, "text": "第1章 性能度量\n1.1 运行时间\n正文。"},
                 {"page": 2, "text": table + "\n1.2 基准评测\n正文。"}]
        self.assertEqual([h["number"] for h in cached_pdf_outline(pages)], ["1", "1.1", "1.2"])
        uncertain = cached_pdf_outline([{"page": 1, "text": "第1章\n模型训练\n1.2 孤立编号\n正文。"}])
        self.assertEqual([entry.get("role") for entry in uncertain], ["boundary"])
        self.assertEqual(cached_pdf_outline([{"page": 1, "text": "1.1 未确认的编号\n正文。\n1.2 另一个编号"}]), [])

    def test_later_repeated_running_headers_do_not_remove_confirmed_sections(self):
        pages = [
            {"page": 1, "text": "第1章 模型训练\n1.1 训练数据\n1.1.1 数据采样\n训练集用于拟合。"},
            {"page": 2, "text": "1.1 训练数据\n2\n正文。"},
            {"page": 3, "text": "1.1 训练数据\n3\n正文。"},
            {"page": 4, "text": "1.1 训练数据\n4\n1.2 模型评估\n测试集用于评估。"},
        ]
        confirmed = cached_pdf_outline(pages[:1])
        self.assertEqual([entry["number"] for entry in confirmed], ["1", "1.1", "1.1.1"])
        for count in (2, 3, 4):
            growing = cached_pdf_outline(pages[:count])
            self.assertEqual(growing[:len(confirmed)], confirmed)
            self.assertEqual(sum(entry["number"] == "1.1" for entry in growing), 1)
        self.assertEqual(cached_pdf_outline(pages)[-1]["number"], "1.2")

    def test_unconfirmed_new_chapter_stops_previous_section_ownership(self):
        pages = [
            {"page": 1, "text": "第1章 模型训练\n1.1 训练数据\n1.1.1 数据采样\n训练集用于拟合。"},
            {"page": 2, "text": "第2章 模型评估\n章首导读。"},
            {"page": 3, "text": "2.1 测试方法\n神经网络使用测试集。"},
            {"page": 4, "text": "第3章 模型部署\n3.1 推理服务\n知识图谱支持关系抽取。"},
        ]
        outline = cached_pdf_outline(pages)
        boundary = next(entry for entry in outline if entry.get("role") == "boundary")
        self.assertEqual((boundary["page"], boundary["number"]), (2, "2"))
        self.assertEqual(pages[1]["text"][boundary["start"]:boundary["end"]], boundary["quote"])
        self.assertNotIn("2.1", {entry["number"] for entry in outline})
        concepts = [node for page in pages for node in extract_local(page["text"], page["page"])["nodes"]]
        tree = hierarchy({"id": "test", "title": "教材", "start_page": 1, "page_kind": "pdf"}, outline, concepts)
        containers = {node["id"]: node for node in tree["nodes"]}
        owners = {m["target"]: containers[m["source"]]["title"] for m in tree["memberships"]}
        self.assertEqual(owners[term_id("训练集")], "1.1.1 数据采样")
        self.assertEqual(owners[term_id("神经网络")], "教材")
        self.assertEqual(owners[term_id("测试集")], "教材")
        self.assertEqual(owners[term_id("知识图谱")], "3.1 推理服务")
        self.assertNotIn("第2章 模型评估", {node["title"] for node in tree["nodes"]})

    def test_assemble_extends_pdf_outline_from_cache_without_reopening_pdf_or_ocr(self):
        ident = self.store.create("cached.pdf")["id"]
        directory = self.store.directory(ident)
        first = {"page": 1, "text": "第1章 模型训练\n1.1 训练数据\n训练集用于拟合。", "ocr": True}
        second = {"page": 2, "text": "1.2 模型评估\n神经网络使用测试集。", "ocr": True}
        atomic_json(directory / "pages/1.json", first)
        atomic_json(directory / "chunks/000001-0000.json", extract_local(first["text"], 1))
        with patch("learning_agent.documents.cached_pdf_outline", wraps=cached_pdf_outline) as infer, \
                patch("fitz.open", side_effect=AssertionError("assemble must use cached pages")), \
                patch("learning_agent.documents.ocr_page", side_effect=AssertionError("no repeat OCR")):
            self.store.assemble(ident)
            before = json.loads((directory / "outline.json").read_text(encoding="utf-8"))
            self.store.assemble(ident)
            self.assertEqual(infer.call_count, 1)
            atomic_json(directory / "pages/2.json", second)
            atomic_json(directory / "chunks/000002-0000.json", extract_local(second["text"], 2))
            graph = self.store.assemble(ident)
            self.assertEqual(infer.call_count, 2)
        after = json.loads((directory / "outline.json").read_text(encoding="utf-8"))
        self.assertEqual(after[:len(before)], before)
        self.assertEqual(after[-1]["title"], "1.2 模型评估")
        self.assertNotIn("page", {n["kind"] for n in graph["hierarchy"]["nodes"]})
        self.assertEqual(self.store.status(ident)["stats"]["ocr_pages"], 2)

    def test_pdf_outline_fallback_preserves_manual_or_embedded_outline(self):
        import fitz
        ident = self.store.create("established.pdf")["id"]
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((40, 60), "Neural networks include neurons. Neural networks learn representations.")
            pdf.save(self.store.source(ident))
        established = [{"level": 1, "title": "Reviewed chapter", "page": 1, "start": 0}]
        path = self.store.directory(ident) / "outline.json"
        atomic_json(path, established)
        with patch("learning_agent.documents.cached_pdf_outline", side_effect=AssertionError("existing outline wins")):
            self.store.run(ident)
        self.assertEqual(json.loads(path.read_text(encoding="utf-8")), established)
        atomic_json(path, [])
        self.store.update(ident, pdf_has_toc=True)
        with patch("learning_agent.documents.cached_pdf_outline", side_effect=AssertionError("embedded TOC exists")):
            self.store.assemble(ident)
        self.assertEqual(json.loads(path.read_text(encoding="utf-8")), [])

    def test_invalid_pdf_and_out_of_range_fail_honestly(self):
        job = self.store.create("broken.pdf")
        self.store.source(job["id"]).write_bytes(b"broken")
        self.store.run(job["id"])
        self.assertEqual(self.store.status(job["id"])["status"], "failed")
        job = self.store.create("text.txt", start_page=10)
        self.store.source(job["id"]).write_text("hello", encoding="utf-8")
        self.store.run(job["id"])
        self.assertIn("范围", self.store.status(job["id"])["error"])

    def test_scanned_pdf_uses_ocr_and_caches_page(self):
        import fitz
        job = self.store.create("scan.pdf")
        with fitz.open() as doc:
            page = doc.new_page()
            pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 50, 50))
            pix.clear_with(255)
            page.insert_image(page.rect, pixmap=pix)
            doc.save(self.store.source(job["id"]))
        with patch("learning_agent.documents.ocr_page", return_value="知识图谱包括实体。神经网络使用训练集。") as ocr:
            self.store.run(job["id"])
            self.store.run(job["id"])
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual(self.store.status(job["id"])["stats"]["ocr_pages"], 1)

    def test_markdown_hierarchy_and_text_sections(self):
        ident = self.book("# 机器学习\n## 知识图谱\n知识图谱包括实体。知识图谱支持关系抽取。", "book.md")
        graph = self.store.graph(ident)
        by_title = {n["title"]: n for n in graph["hierarchy"]["nodes"]}
        self.assertEqual(by_title["知识图谱"]["parent_id"], by_title["机器学习"]["id"])
        self.assertEqual(graph["page_kind"], "section")

    def test_headingless_markdown_connects_00_root_to_every_concept(self):
        ident = self.book("知识图谱包括实体和关系。实体对齐用于知识融合。", "00.md")
        graph = self.store.graph(ident)
        self.assertTrue(graph["nodes"])
        containers = graph["hierarchy"]["nodes"]
        self.assertEqual(len(containers), 1)
        self.assertEqual((containers[0]["title"], containers[0]["kind"]), ("00", "book"))
        memberships = graph["hierarchy"]["memberships"]
        self.assertEqual({m["source"] for m in memberships}, {containers[0]["id"]})
        self.assertEqual({m["target"] for m in memberships}, {n["id"] for n in graph["nodes"]})

    def test_standalone_bold_headings_preserve_original_page_and_offset(self):
        titles = ["P 问题 (Polynomial time，多项式时间)", "NP 问题 (Nondeterministic Polynomial time)", "P 与 NP 的关系"]
        text = "原文介绍。\r\n" + " " * 6000 + "\r\n" + "\r\n正文。\r\n".join(f"**{title}**" for title in titles)
        source = Path(self.tmp.name) / "00.md"
        source.write_bytes(text.encode("utf-8"))
        outline = document_outline(source)
        self.assertEqual([h["title"] for h in outline], titles)
        self.assertEqual([h["level"] for h in outline], [1, 1, 1])
        self.assertTrue(all(h["page"] > 1 for h in outline))
        for heading, title in zip(outline, titles):
            at = (heading["page"] - 1) * 6000 + heading["start"]
            self.assertEqual(at, text.index(f"**{title}**"))
            self.assertTrue(text[at:].startswith(f"**{title}**"))

    def test_markdown_bold_peers_follow_explicit_heading_and_exclude_body_or_code(self):
        text = "\n".join([
            "# 复杂度 ###", "**P 问题**", "正文内容。", "**NP 问题**", "## 判定条件", "**验证过程**",
            "```markdown", "# 代码中的标题", "**代码中的粗体**", "````", "~~~", "### 另一围栏标题", "**另一段代码**", "~~~",
            "- **列表中的粗体**", "1. **有序列表粗体**", "  **列表缩进粗体**", "    **缩进代码**", "> **引用中的粗体**",
            "正文里有**粗体词**。", "**P** 与 **NP**", "**NP 完全问题**，是指一类正文内容。",
            "**" + "这是加粗的长正文段落。" * 20 + "**", "# 总结", "**关键关系**",
        ])
        source = Path(self.tmp.name) / "mixed.md"
        source.write_bytes(text.encode("utf-8"))
        outline = document_outline(source)
        self.assertEqual([(h["level"], h["title"]) for h in outline], [
            (1, "复杂度"), (2, "P 问题"), (2, "NP 问题"), (2, "判定条件"), (3, "验证过程"), (1, "总结"), (2, "关键关系"),
        ])

    def test_bold_section_hierarchy_survives_extraction_and_course_import(self):
        ident = self.book("**P 问题**\n知识图谱包括实体。\n**NP 问题**\n神经网络使用训练集。", "00.md")
        graph = self.store.graph(ident)
        containers = {n["id"]: n for n in graph["hierarchy"]["nodes"]}
        owners = {m["target"]: containers[m["source"]]["title"] for m in graph["hierarchy"]["memberships"]}
        self.assertEqual(owners[term_id("知识图谱")], "P 问题")
        self.assertEqual(owners[term_id("神经网络")], "NP 问题")
        self.assertEqual({n["kind"] for n in containers.values()}, {"book", "chapter"})
        for node in graph["nodes"]:
            for ev in node["evidence"]:
                self.assertEqual(self.store.page(ident, ev["page"])["text"][ev["start"]:ev["end"]], ev["quote"])
        imported = self.store.import_draft(ident, self.course, self.course.load_graph("draft")["version"])["graph"]
        chapter = next(c for c in imported["chapters"] if c["id"] == "doc_" + ident)
        mapped = chapter["document_hierarchy"]
        self.assertEqual(mapped["nodes"], graph["hierarchy"]["nodes"])
        targets = {n["id"]: n for n in imported["nodes"] if n.get("document_id") == ident}
        imported_owners = {targets[m["target"]]["title"]: containers[m["source"]]["title"] for m in mapped["memberships"]}
        self.assertEqual(imported_owners["知识图谱"], "P 问题")
        self.assertEqual(imported_owners["神经网络"], "NP 问题")

    def test_explanation_titles_stay_in_outline_but_are_not_graph_parents(self):
        text = ("**模型训练**\n训练集用于拟合。\n**通俗理解：从例子理解**\n知识图谱包括实体。\n"
                "**严格的判定条件**\n实体对齐用于知识融合。\n**模型评估**\n神经网络使用测试集。")
        ident = self.book(text, "概念说明.md")
        outline = document_outline(self.store.source(ident))
        self.assertEqual(len(outline), 4)
        explanations = [h for h in outline if h.get("role") == "explanation"]
        self.assertEqual([h["title"] for h in explanations], ["通俗理解：从例子理解", "严格的判定条件"])
        original = self.store.page(ident, 1)["text"]
        self.assertTrue(all(original[h["start"]:].startswith(f"**{h['title']}**") for h in explanations))
        graph = self.store.graph(ident)
        containers = {n["id"]: n for n in graph["hierarchy"]["nodes"]}
        self.assertFalse({"通俗理解：从例子理解", "严格的判定条件"} & {n["title"] for n in containers.values()})
        owners = {m["target"]: containers[m["source"]]["title"] for m in graph["hierarchy"]["memberships"]}
        self.assertEqual(owners[term_id("知识图谱")], "模型训练")
        self.assertEqual(owners[term_id("实体对齐")], "模型训练")
        self.assertEqual(owners[term_id("神经网络")], "模型评估")

    def test_text_chunk_boundaries_preserve_proofs_without_creating_chapters(self):
        text = "知识图谱包括实体。".ljust(6050, " ") + "神经网络使用训练集。知识图谱支持关系抽取。"
        ident = self.book(text, "分段材料.txt")
        graph = self.store.graph(ident)
        stats = self.store.status(ident)["stats"]
        self.assertGreater(stats["pages"], 1)
        self.assertGreater(stats["chunks"], stats["pages"])
        self.assertEqual([n["kind"] for n in graph["hierarchy"]["nodes"]], ["book"])
        root = graph["hierarchy"]["nodes"][0]["id"]
        self.assertTrue(all(m["source"] == root for m in graph["hierarchy"]["memberships"]))
        concept = next(n for n in graph["nodes"] if n["title"] == "知识图谱")
        self.assertEqual({ev["page"] for ev in concept["evidence"]}, {1, 2})
        for item in graph["nodes"] + graph["edges"]:
            for ev in item["evidence"]:
                page = self.store.page(ident, ev["page"])
                self.assertEqual(page["text"][ev["start"]:ev["end"]], ev["quote"])
        imported = self.store.import_draft(ident, self.course, self.course.load_graph("draft")["version"])["graph"]
        chapter = next(c for c in imported["chapters"] if c["id"] == "doc_" + ident)
        self.assertEqual(chapter["document_hierarchy"]["nodes"], graph["hierarchy"]["nodes"])
        additions = [n for n in imported["nodes"] if n.get("document_id") == ident]
        self.assertEqual({m["source"] for m in chapter["document_hierarchy"]["memberships"]}, {root})
        self.assertEqual({m["target"] for m in chapter["document_hierarchy"]["memberships"]}, {n["id"] for n in additions})
        originals = {n["title"]: n for n in graph["nodes"]}
        self.assertTrue(all(n["document_evidence"] == originals[n["title"]]["evidence"] for n in additions))

    def test_docx_without_headings_uses_root_across_text_blocks(self):
        from docx import Document
        job = self.store.create("unstructured.docx")
        ident = job["id"]
        document = Document()
        document.add_paragraph("Knowledge graphs link entities.".ljust(6050, " "))
        document.add_paragraph("Neural networks include neurons. Neural networks learn representations.")
        document.save(self.store.source(ident))
        self.store.run(ident)
        graph = self.store.graph(ident)
        self.assertTrue(graph["nodes"])
        self.assertGreater(self.store.status(ident)["stats"]["pages"], 1)
        self.assertEqual([n["kind"] for n in graph["hierarchy"]["nodes"]], ["book"])
        root = graph["hierarchy"]["nodes"][0]["id"]
        self.assertEqual({m["source"] for m in graph["hierarchy"]["memberships"]}, {root})
        for node in graph["nodes"]:
            for ev in node["evidence"]:
                self.assertEqual(self.store.page(ident, ev["page"])["text"][ev["start"]:ev["end"]], ev["quote"])

    def test_pdf_without_outline_uses_root_and_preserves_page_evidence_on_import(self):
        import fitz
        job = self.store.create("无目录.pdf")
        ident = job["id"]
        with fitz.open() as doc:
            for text in ("Neural networks include neurons. Neural networks learn representations.",
                         "Knowledge graphs link entities. Entity alignment connects knowledge graphs."):
                page = doc.new_page()
                page.insert_text((40, 60), text)
            doc.save(self.store.source(ident))
        self.store.run(ident)
        graph = self.store.graph(ident)
        self.assertTrue(graph["nodes"])
        containers = {n["id"]: n for n in graph["hierarchy"]["nodes"]}
        self.assertEqual(len(containers), 1)
        self.assertEqual({n["kind"] for n in containers.values()}, {"book"})
        memberships = graph["hierarchy"]["memberships"]
        self.assertEqual({containers[m["source"]]["kind"] for m in memberships}, {"book"})
        self.assertEqual({ev["page"] for n in graph["nodes"] for ev in n["evidence"]}, {1, 2})
        for item in graph["nodes"] + graph["edges"]:
            for ev in item["evidence"]:
                self.assertEqual(self.store.page(ident, ev["page"])["text"][ev["start"]:ev["end"]], ev["quote"])
        imported = self.store.import_draft(ident, self.course, self.course.load_graph("draft")["version"])["graph"]
        chapter = next(c for c in imported["chapters"] if c["id"] == "doc_" + ident)
        self.assertEqual(chapter["document_hierarchy"]["nodes"], graph["hierarchy"]["nodes"])
        targets = {n["id"] for n in imported["nodes"] if n.get("document_id") == ident}
        self.assertEqual({m["target"] for m in chapter["document_hierarchy"]["memberships"]}, targets)

    def test_names_merge_with_multiple_page_proofs(self):
        ident = self.book("知识图谱包括实体。" + " " * 6000 + "知识图谱支持关系抽取。")
        graph = self.store.graph(ident)
        nodes = [n for n in graph["nodes"] if n["title"] == "知识图谱"]
        self.assertEqual(len(nodes), 1)
        self.assertEqual({ev["page"] for ev in nodes[0]["evidence"]}, {1, 2})

    def test_model_rejects_invented_quotes_and_endpoint_mismatch(self):
        result = {"nodes": [{"id": "a", "title": "知识图谱", "quote": "知识图谱关联神经网络。"}, {"id": "b", "title": "神经网络", "quote": "知识图谱关联神经网络。"}],
                  "edges": [{"source": "a", "target": "b", "type": "contains", "quote": "不存在的原文"}]}
        with patch("learning_agent.llm.call_llm_json", return_value=json.dumps(result)):
            extracted = extract_llm("知识图谱关联神经网络。", 3, 100, {})
        self.assertEqual(len(extracted["edges"]), 0)
        self.assertEqual(extracted["rejected"], 1)
        self.assertEqual(extracted["nodes"][0]["evidence"][0]["start"], 100)
        result["nodes"][0]["quote"] = "虚构"
        result["nodes"][1]["quote"] = "虚构"
        with patch("learning_agent.llm.call_llm_json", return_value=json.dumps(result)):
            with self.assertRaises(CourseGraphError):
                extract_llm("知识图谱关联神经网络。", 1, config={})

    def test_failed_chunk_retry_retains_completed_checkpoints(self):
        job = self.store.create("book.txt")
        self.store.source(job["id"]).write_text("知识图谱包括实体。" * 800, encoding="utf-8")
        real_extract = extract_local
        calls = 0
        def fail_second(*args):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise CourseGraphError("模拟请求失败", 502)
            return real_extract(*args)
        with patch("learning_agent.documents.extract_local", side_effect=fail_second):
            self.store.run(job["id"])
        self.assertEqual(self.store.status(job["id"])["status"], "partial")
        self.assertEqual(self.store.status(job["id"])["stats"]["chunks"], 1)
        first = next((self.store.directory(job["id"]) / "chunks").glob("*.json"))
        before = first.stat().st_mtime_ns
        self.store.run(job["id"])
        self.assertEqual(self.store.status(job["id"])["status"], "completed")
        self.assertEqual(first.stat().st_mtime_ns, before)

    def test_cancel_preserves_status_and_stops_processing(self):
        job = self.store.create("book.txt")
        self.store.source(job["id"]).write_text("知识图谱包括实体。", encoding="utf-8")
        self.store.cancelled.add(job["id"])
        self.store.run(job["id"])
        self.assertEqual(self.store.status(job["id"])["status"], "cancelled")
        self.assertFalse((self.store.directory(job["id"]) / "graph.json").exists())

    def test_import_is_draft_versioned_and_maps_cooccurrence(self):
        ident = self.book()
        before = self.course.load_graph("draft")
        with self.assertRaises(CourseGraphError) as conflict:
            self.store.import_draft(ident, self.course, before["version"] + 1)
        self.assertEqual(conflict.exception.status_code, 409)
        result = self.store.import_draft(ident, self.course, before["version"])
        self.assertEqual(result["version"], before["version"] + 1)
        added = [n for n in result["graph"]["nodes"] if n.get("document_id") == ident]
        self.assertTrue(added)
        self.assertTrue(all(n["review_status"] == "draft" and n["document_evidence"] for n in added))
        self.assertTrue(all(n["document_page_kind"] == "section" for n in added))
        self.assertIsNone(self.course.load_graph("published"))
        for edge in result["graph"]["edges"]:
            if edge.get("extraction_type") == "cooccurs":
                self.assertEqual(edge["type"], "related")
        duplicate = self.store.import_draft(ident, self.course, result["version"])
        self.assertEqual(duplicate["nodes_added"], 0)
        self.assertEqual(duplicate["edges_added"], 0)
        self.assertEqual(duplicate["version"], result["version"])

    def test_incremental_import_adds_missing_nodes_edges_and_mapped_hierarchy(self):
        ident = self.book()
        nodes = self.store.graph(ident)["nodes"]
        version = self.course.load_graph("draft")["version"]
        first = self.store.import_draft(ident, self.course, version, [nodes[0]["id"]])
        second = self.store.import_draft(ident, self.course, first["version"], [n["id"] for n in nodes[1:]])
        self.assertEqual(first["nodes_added"] + second["nodes_added"], len(nodes))
        self.assertEqual(len([s for s in second["graph"]["sources"] if s["id"] == "doc_" + ident]), 1)
        chapter = next(c for c in second["graph"]["chapters"] if c["id"] == "doc_" + ident)
        course_ids = {n["id"] for n in second["graph"]["nodes"]}
        self.assertTrue(all(m["target"] in course_ids for m in chapter["document_hierarchy"]["memberships"]))

    def test_same_page_headings_assign_concepts_by_character_offset(self):
        ident = self.book("# 第一节\n神经网络包括训练集。\n# 第二节\n知识图谱支持实体对齐。", "book.md")
        graph = self.store.graph(ident)
        by_id = {n["id"]: n for n in graph["hierarchy"]["nodes"]}
        membership = {m["target"]: by_id[m["source"]]["title"] for m in graph["hierarchy"]["memberships"]}
        self.assertEqual(membership[term_id("神经网络")], "第一节")
        self.assertEqual(membership[term_id("实体对齐")], "第二节")

    def test_path_traversal_and_invalid_engine_are_rejected(self):
        for value in ["../../secret", "x" * 32, "a" * 33]:
            with self.assertRaises(CourseGraphError):
                self.store.directory(value)
        with self.assertRaises(CourseGraphError):
            self.store.create("book.exe")
        with self.assertRaises(CourseGraphError):
            self.store.create("book.pdf", engine="unknown")
        job = self.store.create("../../secret.txt")
        self.assertEqual(job["filename"], "secret.txt")

    def test_upload_api_and_original_file_round_trip(self):
        from fastapi.testclient import TestClient
        from learning_agent.main import app
        with patch("learning_agent.document_api.document_store", self.store):
            client = TestClient(app)
            response = client.post("/api/documents/upload", files={"file": ("教材.txt", "知识图谱包括实体。关系抽取关联神经网络。".encode("utf-8"), "text/plain")})
            self.assertEqual(response.status_code, 200, response.text)
            ident = response.json()["id"]
            self.store.futures[ident].result(timeout=15)
            self.assertEqual(client.get(f"/api/documents/{ident}").json()["status"], "completed")
            self.assertTrue(client.get(f"/api/documents/{ident}/graph").json()["nodes"])
            self.assertIn("知识图谱", client.get(f"/api/documents/{ident}/source").text)
            self.assertEqual(client.get(f"/api/documents/{ident}/pages/1").status_code, 200)
            self.assertEqual(len(client.get("/api/documents").json()["documents"]), 1)


if __name__ == "__main__":
    unittest.main()
