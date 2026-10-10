"""In-memory product export tests. Not a document layout acceptance test."""
import asyncio
import io
import os
import unittest
import zipfile
from unittest.mock import patch

from learning_agent.course_graph import CourseGraphError
from pliac.word_export import render_word, assemble
from test_personal_export import sample_material
from fastapi.testclient import TestClient
from pliac.main import app
import test_assessment
from pliac.notebook import Notebook
from pliac.reports import StageReports


class WordExportTests(unittest.TestCase):
    def test_nested_lists_preserve_numbering_paragraph_order_and_emphasis(self):
        from lxml import html
        from docx import Document
        source = html.fragment_fromstring('''<ol start="3"><li><p>Parent <strong>important</strong></p>
            <ul><li>First child</li><li>Second child<ol><li>Deep child</li></ol></li></ul>
            <p>Parent continuation</p></li><li>Next parent</li></ol>''', create_parent='div')
        document = Document(io.BytesIO(assemble(source, {})))
        lines = [p.text for p in document.paragraphs]
        self.assertEqual(lines, ['3. Parent important', '3.1. First child', '3.2. Second child',
                                 '3.2.1. Deep child', 'Parent continuation', '4. Next parent'])
        self.assertTrue(next(run for run in document.paragraphs[0].runs if run.text == 'important').bold)
        self.assertGreater(document.paragraphs[3].paragraph_format.left_indent,
                           document.paragraphs[1].paragraph_format.left_indent)
        self.assertEqual(document.paragraphs[4].paragraph_format.left_indent, 0)

    def test_saved_markdown_nested_list_reaches_export_with_all_items(self):
        from docx import Document
        sample = sample_material()
        sample['proposal']['blocks'][0]['text'] = '1. Parent\n    1. Child first\n    2. Child second\n2. Parent second'
        with patch.dict(os.environ, {'PLIAC_WORD_EXPORT': '1'}):
            content = asyncio.run(render_word('material', sample))
        text = [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]
        self.assertIn('1.1. Child first', text)
        self.assertIn('1.2. Child second', text)
        self.assertIn('2. Parent second', text)

    def test_not_enabled_before_layout_verification(self):
        with patch.dict(os.environ, {"PLIAC_WORD_EXPORT": "0"}), self.assertRaises(CourseGraphError) as result:
            asyncio.run(render_word("material", sample_material()))
        self.assertEqual(result.exception.status_code, 503)

    def test_in_memory_package_has_text_fonts_table_and_equation_image(self):
        with patch.dict(os.environ, {"PLIAC_WORD_EXPORT": "1"}):
            content = asyncio.run(render_word("material", sample_material()))
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            xml = package.read("word/document.xml").decode()
            styles = package.read("word/styles.xml").decode()
            self.assertIn("测试集用于最终评价", xml)
            self.assertIn("<w:tbl>", xml)
            self.assertIn("<w:drawing>", xml)
            self.assertNotIn("NOT_FOR_EXPORT", xml)
            self.assertNotIn(r"\frac", xml)
            self.assertIn("华文中宋", styles)
            self.assertIn("宋体", styles)
            self.assertIn("1. ", xml)
            self.assertNotIn('TargetMode="External"', package.read("word/_rels/document.xml.rels").decode())
            self.assertTrue(any(name.startswith("word/media/") for name in package.namelist()))

    def test_invalid_equation_is_not_silently_dropped(self):
        sample = sample_material()
        sample["proposal"]["blocks"][0]["text"] = r"$\notAValidPliacCommand{x}$"
        with patch.dict(os.environ, {"PLIAC_WORD_EXPORT": "1"}), self.assertRaises(CourseGraphError) as result:
            asyncio.run(render_word("material", sample))
        self.assertEqual(result.exception.status_code, 422)


class WordEndpointTests(unittest.TestCase):
    setUp = test_assessment.AssessmentTests.setUp
    payload = test_assessment.AssessmentTests.payload

    def test_saved_report_download_and_missing_owner(self):
        Notebook(self.store).save(self.payload(node_id="a", text="saved snapshot"))
        report = StageReports(self.store).save(self.payload())["workspace"]["stage_reports"][-1]
        with patch("learning_agent.api.store", self.store), patch.dict(os.environ, {"PLIAC_WORD_EXPORT": "1"}), TestClient(app) as client:
            query = {"student_id": "synthetic", "kind": "report", "artifact_id": report["id"]}
            response = client.get("/api/tutor/export/word", params=query)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["cache-control"], "no-store")
            self.assertTrue(response.content.startswith(b"PK"))
            self.assertEqual(client.get("/api/tutor/export/word", params=query | {"student_id": "other"}).status_code, 404)


if __name__ == "__main__":
    unittest.main()
