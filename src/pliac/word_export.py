"""Word export candidate; opt-in until page-render QA is complete.

Uses the same saved-content whitelist as PDF. Equations are rendered images,
not editable OMML. No external links, scripts, or embedded source files.
"""
import io
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import html
from playwright.async_api import async_playwright, Error as BrowserError
from starlette.concurrency import run_in_threadpool

from learning_agent.course_graph import CourseGraphError
from learnmargin.rendering import _math_assets
from .personal_export import export_body


async def equation_images(root):
    equations = root.xpath('.//*[contains(concat(" ", normalize-space(@class), " "), " math ")]')
    if not equations:
        return {}
    css, script = _math_assets()
    result = {}
    async with async_playwright() as runtime:
        browser = await runtime.chromium.launch(headless=True, chromium_sandbox=True)
        try:
            page = await browser.new_page(viewport={"width": 900, "height": 1000}, device_scale_factor=3)
            await page.route("**/*", lambda route: route.abort())
            await page.set_content('<!doctype html><html><head><meta charset="utf-8"><style>' + css + '</style></head><body><span id="formula" style="display:inline-block;font-size:18px;color:black;background:white"></span></body></html>')
            await page.evaluate(script)
            for index, element in enumerate(equations):
                tex = element.get("data-tex", "")
                valid = await page.evaluate('''([tex, display]) => {
                    try {katex.render(tex, document.getElementById('formula'),
                        {displayMode:display, throwOnError:true, trust:false, maxExpand:1000}); return true;}
                    catch (error) {return String(error);}
                }''', [tex, element.get("data-display") == "true"])
                if valid is not True:
                    raise CourseGraphError("公式无法可靠排版，Word 导出未完成，原始材料仍保留。", 422)
                await page.evaluate("document.fonts.ready")
                box = await page.locator("#formula").bounding_box()
                if not box or box["width"] > 640 or box["height"] > 800:
                    raise CourseGraphError("公式超出 Word 正文宽度，请拆分后导出。", 422)
                ident = str(index)
                element.set("data-word-equation", ident)
                result[ident] = (await page.locator("#formula").screenshot(), box["width"] * 25.4 / 96)
        finally:
            await browser.close()
    return result


def assemble(root, equations):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        style = document.styles[name]
        face = "宋体" if name == "Normal" else "华文中宋"
        style.font.name = face
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if attr.endswith("Theme"):
                del fonts.attrib[attr]
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), face)
        style.font.size = Pt(12 if name == "Normal" else 20 if name == "Title" else 15)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.4
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""

    def inline(element, paragraph, bold=False, italic=False):
        ident = element.get("data-word-equation")
        if ident is not None:
            content, width = equations[ident]
            paragraph.add_run().add_picture(io.BytesIO(content), width=Mm(width))
            return
        def text(value):
            if value:
                run = paragraph.add_run(value)
                run.bold, run.italic = bold, italic
        text(element.text)
        for child in element:
            if child.tag == "br":
                paragraph.add_run().add_break()
            else:
                inline(child, paragraph, bold or child.tag in {"strong", "b"}, italic or child.tag in {"em", "i"})
            text(child.tail)

    def numbered_list(element, parent, prefix='', depth=0):
        # Preserve nested structure and source order; bullets become numbered
        # items according to the user's document convention.
        try:
            start = int(element.get('start', '1')) if element.tag == 'ol' else 1
        except ValueError:
            start = 1
        for index, item in enumerate(element, start):
            if item.tag != 'li':
                continue
            number = f'{prefix}{index}'
            paragraph = parent.add_paragraph(number + '. ')
            paragraph.paragraph_format.left_indent = Mm(min(depth, 8) * 6)
            has_content = False
            if item.text and item.text.strip():
                paragraph.add_run(item.text)
                has_content = True
            for child in item:
                if child.tag in {'ul', 'ol'}:
                    numbered_list(child, parent, number + '.', depth + 1)
                    paragraph = None
                else:
                    if paragraph is None or (child.tag == 'p' and has_content):
                        paragraph = parent.add_paragraph()
                        paragraph.paragraph_format.left_indent = Mm(min(depth, 8) * 6)
                    inline(child, paragraph)
                    has_content = True
                if child.tail and child.tail.strip():
                    if paragraph is None:
                        paragraph = parent.add_paragraph()
                        paragraph.paragraph_format.left_indent = Mm(min(depth, 8) * 6)
                    paragraph.add_run(child.tail)
                    has_content = True

    def blocks(element, parent):
        for item in element:
            tag = item.tag
            if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                paragraph = parent.add_paragraph(style="Title" if tag == "h1" else "Heading " + str(int(tag[1]) - 1))
                inline(item, paragraph)
            elif tag in {"p", "pre", "blockquote"} or item.get("data-word-equation") is not None:
                inline(item, parent.add_paragraph())
            elif tag in {"ul", "ol"}:
                numbered_list(item, parent)
            elif tag == "table":
                rows = item.xpath(".//tr")
                columns = max((len(row) for row in rows), default=0)
                if not columns or columns > 8:
                    raise CourseGraphError("表格列数不适合当前 Word 版面，请拆分后导出。", 422)
                table = parent.add_table(rows=0, cols=columns)
                table.autofit = False
                for column in table.columns:
                    column.width = Mm(170 / columns)
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                    border = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "000000")):
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                for row_index, row in enumerate(rows):
                    cells = table.add_row().cells
                    if row_index == 0 and any(cell.tag == "th" for cell in row):
                        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                    for cell, source in zip(cells, row):
                        cell.width = Mm(170 / columns)
                        inline(source, cell.paragraphs[0], source.tag == "th")
            else:
                blocks(item, parent)
    blocks(root, document)
    if equations:
        document.add_paragraph("导出说明：公式以高清图片保留，不是可编辑的 Word 原生公式。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


async def render_word(kind, record):
    if os.environ.get("PLIAC_WORD_EXPORT") != "1":
        raise CourseGraphError("Word 导出尚在版面核验中，请先使用在线保存或 PDF。", 503)
    root = html.fragment_fromstring(export_body(kind, record), create_parent="div")
    try:
        equations = await equation_images(root)
    except BrowserError as error:
        raise CourseGraphError("Word 公式渲染暂不可用，原始材料仍保留。", 503) from error
    return await run_in_threadpool(assemble, root, equations)
